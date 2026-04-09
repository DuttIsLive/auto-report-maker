"""
test_document_formatter.py
--------------------------
Unit tests for document_formatter and report_maker.
"""

from __future__ import annotations

import os
import tempfile

import pytest
from docx import Document

from src.document_formatter import (
    ACCENT_COLOR,
    ReportData,
    Section,
    build_report,
)
from src.report_maker import make_report


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _open_docx(path: str) -> Document:
    return Document(path)


def _all_text(doc: Document) -> str:
    """Collect all text from paragraphs and table cells."""
    parts: list[str] = []
    for para in doc.paragraphs:
        parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    parts.append(para.text)
    return "\n".join(parts)


def _all_fonts(doc: Document) -> set[str]:
    """Return every distinct font name used in the document."""
    fonts: set[str] = set()
    for para in doc.paragraphs:
        for run in para.runs:
            if run.font.name:
                fonts.add(run.font.name)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.font.name:
                            fonts.add(run.font.name)
    return fonts


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestBuildReport:
    def test_creates_file(self, tmp_path):
        data = ReportData(title="Test Report")
        output = str(tmp_path / "out.docx")
        result = build_report(data, output)
        assert result == output
        assert os.path.isfile(output)

    def test_title_appears_in_cover_table(self, tmp_path):
        title = "My Designer Report"
        data = ReportData(title=title)
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        text = _all_text(doc)
        assert title in text

    def test_subtitle_appears(self, tmp_path):
        data = ReportData(title="T", subtitle="My Subtitle")
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        assert "My Subtitle" in _all_text(doc)

    def test_author_appears(self, tmp_path):
        data = ReportData(title="T", author="Jane Doe")
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        assert "Jane Doe" in _all_text(doc)

    def test_date_appears(self, tmp_path):
        data = ReportData(title="T", date="April 2026")
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        assert "April 2026" in _all_text(doc)

    def test_abstract_in_table_cell(self, tmp_path):
        abstract = "This is the abstract text."
        data = ReportData(title="T", abstract=abstract)
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        # Abstract must live inside a table cell, not a plain paragraph
        cell_text = ""
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        cell_text += para.text
        assert abstract in cell_text

    def test_abstract_label_in_table_cell(self, tmp_path):
        data = ReportData(title="T", abstract="Some abstract.")
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        cell_text = ""
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        cell_text += para.text
        assert "Abstract" in cell_text

    def test_sections_appear(self, tmp_path):
        sections = [
            Section(heading="Introduction", body="Intro body text."),
            Section(heading="Methodology",  body="Method body text."),
        ]
        data = ReportData(title="T", sections=sections)
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        text = _all_text(doc)
        assert "Introduction" in text
        assert "Methodology"  in text
        assert "Intro body text." in text
        assert "Method body text." in text

    def test_all_fonts_are_times_new_roman(self, tmp_path):
        data = ReportData(
            title="Font Test",
            subtitle="Sub",
            author="Author",
            date="2026",
            abstract="Abstract text.",
            sections=[Section(heading="Heading", body="Body.")],
        )
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        fonts = _all_fonts(doc)
        assert fonts, "Document has no runs with explicit font names"
        assert fonts == {"Times New Roman"}, (
            f"Expected only Times New Roman, found: {fonts}"
        )

    def test_no_abstract_no_extra_table(self, tmp_path):
        data = ReportData(title="T")
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        # Only the cover-band table should exist (1 table)
        assert len(doc.tables) == 1

    def test_with_abstract_two_tables(self, tmp_path):
        data = ReportData(title="T", abstract="Some abstract.")
        output = str(tmp_path / "out.docx")
        build_report(data, output)
        doc = _open_docx(output)
        # Cover table + abstract box table
        assert len(doc.tables) == 2

    def test_returns_output_path(self, tmp_path):
        data = ReportData(title="T")
        output = str(tmp_path / "report.docx")
        assert build_report(data, output) == output


class TestMakeReport:
    def test_make_report_creates_file(self, tmp_path):
        output = str(tmp_path / "r.docx")
        result = make_report(title="Hello", output_path=output)
        assert result == output
        assert os.path.isfile(output)

    def test_make_report_with_all_fields(self, tmp_path):
        output = str(tmp_path / "full.docx")
        sections = [Section(heading="H1", body="Body 1")]
        make_report(
            title="Full Report",
            output_path=output,
            subtitle="Subtitle",
            author="Author Name",
            date="2026-04-09",
            abstract="Abstract here.",
            sections=sections,
        )
        doc = _open_docx(output)
        text = _all_text(doc)
        assert "Full Report" in text
        assert "Subtitle"    in text
        assert "Author Name" in text
        assert "2026-04-09"  in text
        assert "Abstract here." in text
        assert "H1"          in text
        assert "Body 1"      in text
