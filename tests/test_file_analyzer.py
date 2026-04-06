"""
tests/test_file_analyzer.py
~~~~~~~~~~~~~~~~~~~~~~~~~~~
Unit tests for the file_analyzer module.
"""

import os
import tempfile

import pytest
from docx import Document
from docx.shared import Pt, Inches, Cm

from src.file_analyzer import analyze_files, FileAnalysisResult, _detect_tone, _detect_citation_style


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_docx(paragraphs: list[tuple[str, str]] | None = None,
               margins: dict | None = None) -> str:
    """
    Create a temporary .docx file and return its path.

    Parameters
    ----------
    paragraphs : list of (text, style_name)
    margins : dict with keys top/bottom/left/right in Inches
    """
    doc = Document()
    if margins:
        sec = doc.sections[0]
        sec.top_margin = margins.get("top", Inches(1))
        sec.bottom_margin = margins.get("bottom", Inches(1))
        sec.left_margin = margins.get("left", Inches(1.25))
        sec.right_margin = margins.get("right", Inches(1.25))

    for text, style in (paragraphs or []):
        try:
            p = doc.add_paragraph(text, style=style)
        except KeyError:
            p = doc.add_paragraph(text)

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc.save(tmp.name)
    tmp.close()
    return tmp.name


# ---------------------------------------------------------------------------
# Tests — analyze_files
# ---------------------------------------------------------------------------

def test_analyze_files_empty_list():
    result = analyze_files([])
    assert isinstance(result, FileAnalysisResult)


def test_analyze_files_ignores_non_docx():
    result = analyze_files(["not_a_file.pdf", "also_not.txt"])
    assert isinstance(result, FileAnalysisResult)


def test_analyze_files_single_docx():
    path = _make_docx([
        ("Introduction", "Heading 1"),
        ("This is body text.", "Normal"),
        ("Methods", "Heading 1"),
    ])
    try:
        result = analyze_files([path])
        assert isinstance(result, FileAnalysisResult)
        assert "Introduction" in result.detected_sections
        assert "Methods" in result.detected_sections
    finally:
        os.unlink(path)


def test_analyze_files_extracts_margins():
    path = _make_docx(
        margins={"top": Inches(1.5), "bottom": Inches(1.0),
                 "left": Inches(1.25), "right": Inches(1.25)}
    )
    try:
        result = analyze_files([path])
        assert result.margin_top_cm is not None
        assert result.margin_top_cm > 0
    finally:
        os.unlink(path)


def test_analyze_files_merges_multiple():
    path1 = _make_docx([("Background", "Heading 1"), ("Body text A.", "Normal")])
    path2 = _make_docx([("Results", "Heading 1"), ("Body text B.", "Normal")])
    try:
        result = analyze_files([path1, path2])
        all_sections = result.detected_sections
        assert "Background" in all_sections
        assert "Results" in all_sections
    finally:
        os.unlink(path1)
        os.unlink(path2)


def test_analyze_files_deduplicates_sections():
    path1 = _make_docx([("Introduction", "Heading 1")])
    path2 = _make_docx([("Introduction", "Heading 1")])
    try:
        result = analyze_files([path1, path2])
        count = result.detected_sections.count("Introduction")
        assert count == 1
    finally:
        os.unlink(path1)
        os.unlink(path2)


def test_analyze_files_nonexistent_file():
    """Non-existent files should not raise; they are silently skipped."""
    result = analyze_files(["/tmp/this_does_not_exist_12345.docx"])
    assert isinstance(result, FileAnalysisResult)


# ---------------------------------------------------------------------------
# Tests — internal helpers
# ---------------------------------------------------------------------------

def test_detect_tone_academic():
    text = "This study presents a methodology for analysing research findings and literature."
    assert _detect_tone(text) == "academic"


def test_detect_tone_technical():
    text = "The system architecture includes a deployment pipeline and algorithm configuration."
    assert _detect_tone(text) == "technical"


def test_detect_tone_business():
    text = "The executive quarterly ROI forecast shows positive market strategy outcomes."
    assert _detect_tone(text) == "business"


def test_detect_tone_unknown_returns_professional():
    text = "The weather was fine today."
    assert _detect_tone(text) == "professional"


def test_detect_citation_style_apa():
    text = "Smith (2023). doi:10.1000/xyz. retrieved from Journal of Science."
    result = _detect_citation_style(text)
    assert result == "APA"


def test_detect_citation_style_ieee():
    text = "References [1] [2] [3] IEEE proceedings."
    result = _detect_citation_style(text)
    assert result == "IEEE"


def test_detect_citation_style_none():
    text = "Some random text without clear citation markers."
    result = _detect_citation_style(text)
    assert result is None


# ---------------------------------------------------------------------------
# Tests — FileAnalysisResult.summary()
# ---------------------------------------------------------------------------

def test_file_analysis_result_summary():
    result = FileAnalysisResult(
        body_font="Times New Roman",
        heading_font="Times New Roman",
        body_font_size_pt=12.0,
        detected_tone="academic",
        detected_sections=["Introduction", "Conclusion"],
    )
    summary = result.summary()
    assert "Times New Roman" in summary
    assert "academic" in summary
    assert "Introduction" in summary
