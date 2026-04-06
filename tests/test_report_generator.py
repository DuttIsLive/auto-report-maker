"""
tests/test_report_generator.py
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
Integration tests for the report_generator orchestration pipeline.
"""

import os
import tempfile

import pytest
from docx import Document

from src.report_generator import generate_report


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _tmp_output() -> str:
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    return tmp.name


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

def test_generate_report_produces_docx(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    out = _tmp_output()
    try:
        result = generate_report(
            topic="Artificial Intelligence",
            output_path=out,
            use_ai=False,
            verbose=False,
        )
        assert os.path.isfile(result)
        assert result.endswith(".docx")
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_generate_report_valid_word_document(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    out = _tmp_output()
    try:
        result = generate_report(
            topic="Climate Science",
            output_path=out,
            use_ai=False,
            verbose=False,
        )
        doc = Document(result)
        assert len(doc.paragraphs) > 0
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_generate_report_default_output_path(monkeypatch, tmp_path, monkeypatch_cwd=None):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.chdir(tmp_path)
    result = generate_report(
        topic="Neural Networks",
        output_path=None,
        use_ai=False,
        verbose=False,
    )
    assert os.path.isfile(result)
    assert "Neural_Networks" in os.path.basename(result)
    os.unlink(result)


def test_generate_report_raises_on_empty_topic(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    with pytest.raises(ValueError, match="topic must not be empty"):
        generate_report(topic="  ", use_ai=False, verbose=False)


def test_generate_report_with_num_pages(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    out = _tmp_output()
    try:
        result = generate_report(
            topic="Quantum Computing",
            output_path=out,
            num_pages=8,
            use_ai=False,
            verbose=False,
        )
        assert os.path.isfile(result)
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_generate_report_with_template_files(monkeypatch, tmp_path):
    """Template files that exist are analysed; missing ones are skipped gracefully."""
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    out = _tmp_output()
    # Create a minimal valid .docx template
    from docx import Document as _Doc
    tpl_path = str(tmp_path / "template.docx")
    d = _Doc()
    d.add_heading("Introduction", 1)
    d.add_paragraph("Some body text for tone detection.")
    d.save(tpl_path)
    try:
        result = generate_report(
            topic="Robotics",
            output_path=out,
            file_paths=[tpl_path],
            use_ai=False,
            verbose=False,
        )
        assert os.path.isfile(result)
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_generate_report_verbose_prints(monkeypatch, capsys):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    out = _tmp_output()
    try:
        generate_report(
            topic="Space Exploration",
            output_path=out,
            use_ai=False,
            verbose=True,
        )
        captured = capsys.readouterr()
        assert "Space Exploration" in captured.out
        assert "Report saved" in captured.out
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_generate_report_document_contains_topic(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    out = _tmp_output()
    try:
        result = generate_report(
            topic="Sustainable Energy",
            output_path=out,
            use_ai=False,
            verbose=False,
        )
        doc = Document(result)
        full_text = " ".join(p.text for p in doc.paragraphs)
        assert "Sustainable Energy" in full_text
    finally:
        if os.path.exists(out):
            os.unlink(out)
